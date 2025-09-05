#!/usr/bin/env python3
"""
파일 내용을 텍스트로 추출하는 유틸리티 클래스
다양한 파일 형식 지원: TXT, MD, PDF, DOCX, PPTX, XLSX 등
"""

import os
import mimetypes
from pathlib import Path
from typing import Optional, Tuple
import base64
from io import BytesIO

# 조건부 import - 라이브러리가 없어도 기본 기능은 동작하도록
try:
    from markitdown import MarkItDown
    MARKITDOWN_AVAILABLE = True
except ImportError:
    MARKITDOWN_AVAILABLE = False

try:
    import PyPDF2
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from PIL import Image
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False


class FileExtractor:
    """다양한 파일 형식에서 텍스트를 추출하는 클래스"""
    
    def __init__(self):
        self.markitdown = None
        if MARKITDOWN_AVAILABLE:
            try:
                self.markitdown = MarkItDown()
            except Exception:
                pass
    
    def get_supported_extensions(self) -> list:
        """지원하는 파일 확장자 목록 반환"""
        extensions = ['.txt', '.md', '.py', '.js', '.html', '.css', '.json', 
                     '.xml', '.yml', '.yaml', '.ini', '.conf', '.log']
        
        if PYPDF2_AVAILABLE or MARKITDOWN_AVAILABLE:
            extensions.append('.pdf')
        
        if DOCX_AVAILABLE or MARKITDOWN_AVAILABLE:
            extensions.append('.docx')
        
        if OPENPYXL_AVAILABLE or MARKITDOWN_AVAILABLE:
            extensions.append('.xlsx')
        
        if PPTX_AVAILABLE or MARKITDOWN_AVAILABLE:
            extensions.append('.pptx')
        
        # 이미지 형식 추가
        if PIL_AVAILABLE:
            extensions.extend(['.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp', '.ico'])
        
        return extensions
    
    def extract_text(self, file_path: str) -> Tuple[bool, str, Optional[str]]:
        """
        파일에서 텍스트를 추출
        
        Args:
            file_path: 추출할 파일 경로
            
        Returns:
            Tuple[success: bool, content: str, error_message: Optional[str]]
        """
        try:
            if not os.path.exists(file_path):
                return False, "", "파일이 존재하지 않습니다"
            
            # 파일 크기 확인 (10MB 제한)
            file_size = os.path.getsize(file_path)
            if file_size > 10 * 1024 * 1024:  # 10MB
                return False, "", "파일이 너무 큽니다 (최대 10MB)"
            
            file_path_obj = Path(file_path)
            extension = file_path_obj.suffix.lower()
            
            # markitdown 우선 시도 (가장 포괄적)
            if self.markitdown:
                try:
                    result = self.markitdown.convert(file_path)
                    if result and hasattr(result, 'text_content'):
                        content = result.text_content.strip()
                        if content:
                            return True, content, None
                except Exception as e:
                    # markitdown 실패시 개별 라이브러리로 폴백
                    pass
            
            # 개별 형식별 처리
            if extension in ['.txt', '.md', '.py', '.js', '.html', '.css', 
                           '.json', '.xml', '.yml', '.yaml', '.ini', '.conf', '.log']:
                return self._extract_text_file(file_path)
            
            elif extension == '.pdf':
                return self._extract_pdf(file_path)
            
            elif extension == '.docx':
                return self._extract_docx(file_path)
            
            elif extension == '.xlsx':
                return self._extract_xlsx(file_path)
            
            elif extension == '.pptx':
                return self._extract_pptx(file_path)
            
            else:
                # 알 수 없는 형식은 텍스트 파일로 시도
                return self._extract_text_file(file_path)
                
        except Exception as e:
            return False, "", f"파일 처리 중 오류: {str(e)}"
    
    def _extract_text_file(self, file_path: str) -> Tuple[bool, str, Optional[str]]:
        """일반 텍스트 파일 추출"""
        try:
            encodings = ['utf-8', 'utf-8-sig', 'cp949', 'euc-kr', 'latin-1']
            
            for encoding in encodings:
                try:
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read().strip()
                        return True, content, None
                except UnicodeDecodeError:
                    continue
            
            return False, "", "지원되지 않는 문자 인코딩입니다"
            
        except Exception as e:
            return False, "", f"텍스트 파일 읽기 오류: {str(e)}"
    
    def _extract_pdf(self, file_path: str) -> Tuple[bool, str, Optional[str]]:
        """PDF 파일에서 텍스트 추출"""
        if not PYPDF2_AVAILABLE:
            return False, "", "PDF 처리 라이브러리가 설치되지 않았습니다"
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                text = ""
                
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    text += page.extract_text() + "\n"
                
                return True, text.strip(), None
                
        except Exception as e:
            return False, "", f"PDF 처리 오류: {str(e)}"
    
    def _extract_docx(self, file_path: str) -> Tuple[bool, str, Optional[str]]:
        """DOCX 파일에서 텍스트 추출"""
        if not DOCX_AVAILABLE:
            return False, "", "DOCX 처리 라이브러리가 설치되지 않았습니다"
        
        try:
            doc = Document(file_path)
            text = ""
            
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # 테이블 내용도 추출
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + "\t"
                    text += "\n"
            
            return True, text.strip(), None
            
        except Exception as e:
            return False, "", f"DOCX 처리 오류: {str(e)}"
    
    def _extract_xlsx(self, file_path: str) -> Tuple[bool, str, Optional[str]]:
        """XLSX 파일에서 텍스트 추출"""
        if not OPENPYXL_AVAILABLE:
            return False, "", "XLSX 처리 라이브러리가 설치되지 않았습니다"
        
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            text = ""
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text += f"\n--- {sheet_name} ---\n"
                
                for row in sheet.iter_rows(values_only=True):
                    row_text = []
                    for cell in row:
                        if cell is not None:
                            row_text.append(str(cell))
                        else:
                            row_text.append("")
                    text += "\t".join(row_text) + "\n"
            
            return True, text.strip(), None
            
        except Exception as e:
            return False, "", f"XLSX 처리 오류: {str(e)}"
    
    def _extract_pptx(self, file_path: str) -> Tuple[bool, str, Optional[str]]:
        """PPTX 파일에서 텍스트 추출"""
        if not PPTX_AVAILABLE:
            return False, "", "PPTX 처리 라이브러리가 설치되지 않았습니다"
        
        try:
            presentation = Presentation(file_path)
            text = ""
            
            for slide_num, slide in enumerate(presentation.slides, 1):
                text += f"\n--- Slide {slide_num} ---\n"
                
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
            
            return True, text.strip(), None
            
        except Exception as e:
            return False, "", f"PPTX 처리 오류: {str(e)}"
    
    def get_file_info(self, file_path: str) -> dict:
        """파일 정보 반환"""
        try:
            file_path_obj = Path(file_path)
            file_size = os.path.getsize(file_path)
            extension = file_path_obj.suffix.lower()
            
            # 이미지 파일인지 확인
            is_image = extension in ['.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp', '.ico']
            
            return {
                'name': file_path_obj.name,
                'extension': extension,
                'size': file_size,
                'size_mb': round(file_size / (1024 * 1024), 2),
                'supported': extension in self.get_supported_extensions(),
                'is_image': is_image
            }
        except Exception:
            return {
                'name': 'Unknown',
                'extension': '',
                'size': 0,
                'size_mb': 0,
                'supported': False,
                'is_image': False
            }
    
    def is_image_file(self, file_path: str) -> bool:
        """파일이 이미지인지 확인"""
        extension = Path(file_path).suffix.lower()
        return extension in ['.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp', '.ico']
    
    def process_image(self, file_path: str, max_size: int = 1024) -> Tuple[bool, str, Optional[str]]:
        """
        이미지를 처리하여 base64로 인코딩
        
        Args:
            file_path: 이미지 파일 경로
            max_size: 최대 크기 (가로 또는 세로)
            
        Returns:
            Tuple[success: bool, base64_data: str, error_message: Optional[str]]
        """
        if not PIL_AVAILABLE:
            return False, "", "PIL library is not available"
        
        try:
            # 이미지 열기
            with Image.open(file_path) as img:
                # RGBA를 RGB로 변환 (PNG 투명도 처리)
                if img.mode in ('RGBA', 'LA'):
                    background = Image.new('RGB', img.size, (255, 255, 255))
                    background.paste(img, mask=img.split()[-1] if img.mode == 'RGBA' else img.split()[1])
                    img = background
                elif img.mode not in ('RGB', 'L'):
                    img = img.convert('RGB')
                
                # 이미지 크기 조정
                img.thumbnail((max_size, max_size), Image.Resampling.LANCZOS)
                
                # BytesIO에 저장
                buffer = BytesIO()
                img.save(buffer, format='JPEG', quality=85, optimize=True)
                buffer.seek(0)
                
                # base64로 인코딩
                img_base64 = base64.b64encode(buffer.read()).decode('utf-8')
                
                # 이미지 형식과 함께 data URL 생성
                data_url = f"data:image/jpeg;base64,{img_base64}"
                
                return True, data_url, None
                
        except Exception as e:
            return False, "", f"Error processing image: {str(e)}"
